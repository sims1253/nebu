"""Reference-data adapters for JSON, CSV, and XLSX."""

from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Literal

from openpyxl import load_workbook

# Streaming cap so a small but densely packed workbook cannot expand into
# gigabytes of rows before the pipeline ever starts.
_MAX_XLSX_CELLS = 5_000_000


class ReferenceDataError(ValueError):
    pass


@dataclass(frozen=True)
class ReferenceSource:
    content: bytes
    filename: str

    @property
    def format(self) -> Literal["json", "csv", "xlsx"]:
        suffix = Path(self.filename).suffix.lower().lstrip(".")
        if suffix not in {"json", "csv", "xlsx"}:
            raise ReferenceDataError("reference filename must end in .json, .csv, or .xlsx")
        return suffix  # type: ignore[return-value]

    @classmethod
    def from_path(cls, path: Path) -> ReferenceSource:
        return cls(content=path.read_bytes(), filename=path.name)


def _require_unique_headers(headers: list[str], origin: str) -> None:
    normalized = [value.strip() if value is not None else "" for value in headers]
    if not all(normalized) or len(normalized) != len(set(normalized)):
        raise ReferenceDataError(f"{origin} must have non-empty, unique column headers")


def normalize_reference(
    source: ReferenceSource,
) -> tuple[Any, Literal["json", "csv", "xlsx"]]:
    try:
        if source.format == "json":
            return json.loads(source.content.decode("utf-8-sig")), source.format
        if source.format == "csv":
            text = source.content.decode("utf-8-sig")
            header_line = next(csv.reader(io.StringIO(text)), None)
            if header_line is not None:
                _require_unique_headers(header_line, "CSV reference")
            rows = list(csv.DictReader(io.StringIO(text)))
            tree: dict[str, Any] = {"rows": rows}
            # A two-column CSV is a key/value table in practice; expose a
            # keyed view so specifications bind fields by identity instead of
            # row position. Later duplicate keys win.
            if header_line is not None and len(header_line) == 2:
                key, value = header_line
                tree["lookup"] = {
                    str(row[key]): {"value": row[value]} for row in rows if row.get(key)
                }
            return tree, source.format

        workbook = load_workbook(io.BytesIO(source.content), read_only=True, data_only=True)
        sheets: dict[str, dict[str, list[dict[str, Any]]]] = {}
        cells_seen = 0
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            header_row = next(rows, None)
            if header_row is None:
                sheets[worksheet.title] = {"rows": []}
                continue
            headers = [str(value).strip() if value is not None else "" for value in header_row]
            _require_unique_headers(headers, f"sheet {worksheet.title!r}")
            sheet_rows: list[dict[str, Any]] = []
            for row in rows:
                cells_seen += len(row)
                if cells_seen > _MAX_XLSX_CELLS:
                    raise ReferenceDataError(
                        f"xlsx reference exceeds the {_MAX_XLSX_CELLS:,}-cell limit"
                    )
                if any(value is not None for value in row):
                    sheet_rows.append(dict(zip(headers, row, strict=False)))
            sheets[worksheet.title] = {"rows": sheet_rows}
        return {"sheets": sheets}, source.format
    except ReferenceDataError:
        raise
    except Exception as exc:
        # Any parser failure (bad zip container, corrupt records, encoding
        # problems, ...) is an input problem, not a server error.
        raise ReferenceDataError(f"could not read {source.format} reference data: {exc}") from exc


def resolve_json_pointer(document: Any, pointer: str) -> tuple[bool, Any]:
    """Resolve RFC 6901 JSON Pointer and distinguish missing from explicit null."""
    if pointer == "":
        return True, document
    if not pointer.startswith("/"):
        return False, None
    current = document
    for raw_token in pointer[1:].split("/"):
        token = raw_token.replace("~1", "/").replace("~0", "~")
        if isinstance(current, dict):
            if token not in current:
                return False, None
            current = current[token]
        elif isinstance(current, list):
            # RFC 6901 array tokens are ASCII digits without leading zeros;
            # str.isdigit() alone also accepts superscripts, which int() rejects.
            if not token.isascii() or not token.isdigit() or (len(token) > 1 and token[0] == "0"):
                return False, None
            index = int(token)
            if index >= len(current):
                return False, None
            current = current[index]
        else:
            return False, None
    return True, current
